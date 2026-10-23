from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.export import safe_cell
from app.models import ConsentRecord, PatternDataset, ResearchParticipant, ResearchResponse

EXPECTED_SHEETS = ["Analysis_Records", "Source_Results", "Research_Responses", "Pattern_Dataset", "Consent_Records", "Summary"]


def _count(db, model):
    return db.scalar(select(func.count()).select_from(model))


def _full_flow(client):
    s = client.post("/api/research/start", json={"consent_research": True, "consent_secondary": False}).json()
    rid, pre, post = s["research_id"], s["pre_case"], s["post_case"]
    assert pre["case_id"] != post["case_id"] and "verified_label" not in pre
    base = {"research_id": rid, "decision_reason_code": "link", "response_time_ms": 5000}
    assert client.post("/api/research/answer", json={**base, "case_id": pre["case_id"], "phase": "pre", "user_decision": "scam"}).status_code == 200
    # 未完成初步判斷前不能看輔助資訊
    assert client.post("/api/research/answer", json={**base, "case_id": post["case_id"], "phase": "post_initial", "user_decision": "legit"}).status_code == 200
    help_ = client.post("/api/research/assistance", json={"research_id": rid}).json()
    assert help_["type"] in ("general_tips", "system_evidence")
    r = client.post("/api/research/answer", json={**base, "case_id": post["case_id"], "phase": "post_assisted", "user_decision": "scam"})
    assert r.status_code == 200 and len(r.json()["reveal"]) == 2
    return rid


def test_research_consent_declined(client, db):
    r = client.post("/api/research/start", json={"consent_research": False})
    assert r.status_code == 400 and "基本查核服務不受影響" in r.json()["detail"]
    assert _count(db, ResearchParticipant) == 0 and _count(db, ConsentRecord) == 0
    # 拒絕研究同意仍可使用查核
    assert client.post("/api/analyze/text", json={"text": "請立即匯款到這個帳戶"}).status_code == 200


def test_research_flow_order_and_duplicates(client, db):
    s = client.post("/api/research/start", json={"consent_research": True}).json()
    base = {"research_id": s["research_id"], "decision_reason_code": "link", "response_time_ms": 100}
    early = client.post("/api/research/answer", json={**base, "case_id": s["post_case"]["case_id"], "phase": "post_initial", "user_decision": "scam"})
    assert early.status_code == 409
    ok = {**base, "case_id": s["pre_case"]["case_id"], "phase": "pre", "user_decision": "scam"}
    assert client.post("/api/research/answer", json=ok).status_code == 200
    assert client.post("/api/research/answer", json=ok).status_code == 409


def test_research_consent_withdraw_and_delete(client, db):
    rid = _full_flow(client)
    assert _count(db, ResearchResponse) == 3
    r = client.post("/api/consent", json={"research_id": rid, "purpose": "research", "status": "withdrawn"})
    assert r.status_code == 200
    assert client.post("/api/research/assistance", json={"research_id": rid}).status_code == 404
    d = client.post("/api/privacy/delete", json={"research_id": rid}).json()
    assert d["deleted"]["research_responses"] == 3
    assert _count(db, ResearchResponse) == 0
    statuses = {c.consent_status for c in db.scalars(select(ConsentRecord).where(ConsentRecord.research_id == rid))}
    assert statuses <= {"withdrawn", "declined"}


def test_consented_analysis_linked_to_research(client, db):
    rid = client.post("/api/consent", json={"purpose": "research", "status": "granted"}).json()["research_id"]
    body = client.post("/api/analyze/text", json={"text": "請立即匯款到這個帳戶", "research_id": rid}).json()
    assert body["consent_research"] is True
    client.post("/api/consent", json={"research_id": rid, "purpose": "research", "status": "withdrawn"})
    body2 = client.post("/api/analyze/text", json={"text": "請立即匯款到這個帳戶", "research_id": rid}).json()
    assert body2["consent_research"] is False  # 撤回後不再連結研究資料


def test_export_requires_auth(client):
    assert client.get("/api/admin/export.xlsx").status_code == 401
    assert client.get("/api/admin/export.xlsx", headers={"Authorization": "Bearer forged"}).status_code == 401
    assert client.get("/api/admin/statistics").status_code == 401
    assert client.post("/api/admin/login", json={"username": "admin", "password": "wrong"}).status_code == 401


def test_admin_not_configured(client, settings):
    settings.admin_password_hash = ""
    assert client.post("/api/admin/login", json={"username": "admin", "password": "x"}).status_code == 503


def test_excel_opens_with_expected_structure(client, admin_headers):
    _full_flow(client)
    r = client.get("/api/admin/export.xlsx?dataset=demo", headers=admin_headers)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(BytesIO(r.content))
    assert all(name in wb.sheetnames for name in EXPECTED_SHEETS)
    ws = wb["Research_Responses"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws["B1"].value == "作答編號"
    assert ws.max_row == 4 and ws["A2"].value == "DEMO"
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["資料類型"] == "模擬資料（DEMO）" and summary["有效樣本數（完成前後測）"] == 1


def test_excel_formula_injection(client, db, admin_headers):
    assert safe_cell("=1+1") == "'=1+1" and safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("正常") == "正常"
    db.add(PatternDataset(scam_category='=HYPERLINK("http://evil","x")', psychological_trigger="+cmd", risk_indicators="-1",
                          prevention_method="ok", evidence_source="t", source_license="t"))
    db.commit()
    wb = load_workbook(BytesIO(client.get("/api/admin/export.xlsx", headers=admin_headers).content))
    ws = wb["Pattern_Dataset"]
    cells = [c for row in ws.iter_rows(min_row=2) for c in row]
    assert all(c.data_type != "f" for c in cells)
    assert any(c.value == "'=HYPERLINK(\"http://evil\",\"x\")" for c in cells)


def test_empty_database_export(client, admin_headers):
    st = client.get("/api/admin/statistics", headers=admin_headers).json()
    assert st["valid_samples"] == 0 and st["judgement"]["accuracy"] is None and st["risk_tag_counts"] == {}
    wb = load_workbook(BytesIO(client.get("/api/admin/export.xlsx", headers=admin_headers).content))
    assert wb["Research_Responses"].max_row == 1  # 只有標題列，不產生假資料
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows(min_row=2)}
    assert summary["有效樣本數（完成前後測）"] == 0 and summary["判斷正確率"] == "無資料"


def test_demo_and_formal_data_isolated(client, db, admin_headers, settings):
    assert client.post("/api/admin/demo/seed", headers=admin_headers).json()["participants"] == 12
    demo = client.get("/api/admin/statistics?dataset=demo", headers=admin_headers).json()
    formal = client.get("/api/admin/statistics?dataset=formal", headers=admin_headers).json()
    assert demo["participants"] == 12 and formal["participants"] == 0 and formal["valid_samples"] == 0
    # 切換為正式模式後（需先核准案例），新資料不會被標成 DEMO
    settings.research_mode = "formal"
    assert client.post("/api/research/start", json={"consent_research": True}).status_code == 503  # 無已核准案例
    cases = client.get("/api/admin/cases", headers=admin_headers).json()
    for c in cases:
        client.patch(f"/api/admin/cases/{c['case_id']}", json={"status": "approved"}, headers=admin_headers)
    s = client.post("/api/research/start", json={"consent_research": True}).json()
    assert s["is_demo"] is False
    assert client.get("/api/admin/statistics?dataset=formal", headers=admin_headers).json()["participants"] == 1
    # 清除測試資料只刪除 DEMO
    client.post("/api/admin/demo/clear", headers=admin_headers)
    assert client.get("/api/admin/statistics?dataset=demo", headers=admin_headers).json()["participants"] == 0
    assert client.get("/api/admin/statistics?dataset=formal", headers=admin_headers).json()["participants"] == 1


def test_research_collection_stopped(client, settings):
    settings.research_collection_enabled = False
    assert client.post("/api/research/start", json={"consent_research": True}).status_code == 403


def test_feedback_and_survey(client, admin_headers):
    aid = client.post("/api/analyze/text", json={"text": "請立即匯款到這個帳戶"}).json()["analysis_id"]
    assert client.post("/api/feedback", json={"analysis_id": aid, "helpful": True, "rating": 4}).status_code == 200
    assert client.post("/api/feedback", json={"analysis_id": "0" * 32}).status_code == 404
    survey = {"respondent_role": "teacher", "product": "B", "perceived_value": 4, "most_needed_feature": "teaching_material",
              "teacher_would_use": True, "verbal_interest": True, "trial_intent": True, "payment_intent": False, "acceptable_price_range": "1-99"}
    assert client.post("/api/survey", json=survey).status_code == 200
    client.post("/api/admin/business-events", json={"product": "B", "event_type": "actual_trial", "count": 2}, headers=admin_headers)
    b = client.get("/api/admin/statistics?dataset=demo", headers=admin_headers).json()["business"]["B"]
    assert b["trial_intent"] == 1 and b["actual_trial"] == 2 and b["actual_payment"] == 0


def test_first_run_setup(client, settings, monkeypatch, tmp_path):
    from app import config
    from app.api import admin as admin_api
    monkeypatch.setattr(config, "CONFIG_FILE", tmp_path / "config.env")
    settings.admin_password_hash = ""
    assert client.get("/api/admin/setup-status").json() == {"configured": False, "setup_allowed": False}
    body = {"username": "teacher", "password": "long-enough-pw"}
    # 非本機連線不能設定
    assert client.post("/api/admin/setup", json=body).status_code == 403
    monkeypatch.setattr(admin_api, "_is_local", lambda request: True)
    assert client.post("/api/admin/setup", json={"username": "t", "password": "short"}).status_code == 422
    r = client.post("/api/admin/setup", json=body)
    assert r.status_code == 200 and r.json()["token"]
    saved = (tmp_path / "config.env").read_text(encoding="utf-8")
    assert "ADMIN_PASSWORD_HASH='pbkdf2_sha256$" in saved and "long-enough-pw" not in saved
    assert client.post("/api/admin/setup", json=body).status_code == 409
    assert client.post("/api/admin/login", json=body).status_code == 200
    # 寫入的設定檔可被重新讀取
    from app.config import Settings
    assert Settings(_env_file=str(tmp_path / "config.env")).admin_password_hash.startswith("pbkdf2_sha256$")


def test_proxy_aware_client_ip():
    from types import SimpleNamespace
    from app.security import client_ip, is_local_request

    def req(peer, headers=None):
        return SimpleNamespace(client=SimpleNamespace(host=peer), headers={k.lower(): v for k, v in (headers or {}).items()})
    # 本機直接操作
    assert is_local_request(req("127.0.0.1")) and client_ip(req("127.0.0.1")) == "127.0.0.1"
    # 經 Cloudflare Tunnel 轉進來的外部訪客：不是本機，IP 用 CF-Connecting-IP
    tunneled = req("127.0.0.1", {"CF-Connecting-IP": "203.0.113.9"})
    assert not is_local_request(tunneled) and client_ip(tunneled) == "203.0.113.9"
    # 區網裝置偽造標頭無效
    spoof = req("192.168.1.20", {"CF-Connecting-IP": "1.2.3.4"})
    assert client_ip(spoof) == "192.168.1.20" and not is_local_request(spoof)


def test_tunnel_visitor_cannot_run_first_setup(client, settings, monkeypatch, tmp_path):
    from app import config
    from app.api import admin as admin_api
    from app.security import is_local_request
    monkeypatch.setattr(config, "CONFIG_FILE", tmp_path / "config.env")
    settings.admin_password_hash = ""
    # 讓 TestClient 看起來像本機連線，但帶有 Tunnel 標頭
    monkeypatch.setattr(admin_api, "_is_local", lambda r: is_local_request(type("R", (), {"client": type("C", (), {"host": "127.0.0.1"}), "headers": r.headers})))
    h = {"CF-Connecting-IP": "203.0.113.9"}
    assert client.get("/api/admin/setup-status", headers=h).json()["setup_allowed"] is False
    assert client.post("/api/admin/setup", json={"username": "evil", "password": "x" * 12}, headers=h).status_code == 403
    assert client.get("/api/admin/setup-status").json()["setup_allowed"] is True
